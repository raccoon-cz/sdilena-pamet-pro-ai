"""End-to-end běh na syntetických datech, bez sítě."""
import geopandas as gpd
import openpyxl
import pyogrio

from parcely import testdata as td
from parcely.config import build_config
from parcely.links import mapy_url, nahlizeni_url
from parcely.pipeline import format_summary, run


def make_cfg(tmp_path, extra_ku=None):
    td.create_test_data(tmp_path / "data")
    kus = [{"kod": td.TEST_KU, "nazev": "Testov"}] + (extra_ku or [])
    raw = {
        "obce": [{
            "nazev": "Testov",
            "uzemni_plan": {"cesta": str(tmp_path / "data/up/testov/PlochyRZV_p.shp"), "atribut_kod": "Typ"},
            "katastralni_uzemi": kus,
        }],
        "cache_dir": str(tmp_path / "data/cache"),
        "output_dir": str(tmp_path / "out"),
        # jakýkoli pokus o síť selže okamžitě (nic se nesmí stahovat)
        "zdroje": {"shp_url": "http://127.0.0.1:9/{kod_ku}.zip", "pokusy": 1, "backoff_s": 0},
    }
    return build_config(raw, None)


def test_cely_beh(tmp_path):
    cfg = make_cfg(tmp_path)
    res = run(cfg, out_dir=tmp_path / "out")
    assert res.ok
    ids = {str(td.parcel_id(p)) for p in td.EXPECTED_CANDIDATES}
    assert set(res.kandidati["id_parcely"]) == ids
    assert res.kandidati["skore"].is_monotonic_decreasing
    lv = res.lesni_vyrazene.set_index("id_parcely")
    assert {i for i, d in lv["duvod"].items() if d == "lesní pozemek"} == {str(td.parcel_id("N8"))}
    assert set(lv.index) == {str(td.parcel_id(p)) for p in td.EXPECTED_FOREST | td.EXPECTED_REJECTED}

    k = res.kandidati.set_index("id_parcely")
    n1 = k.loc[str(td.parcel_id("N1"))]
    assert n1["skore"] == 100 and n1["pristup_txt"] == "ano"
    assert k.loc[str(td.parcel_id("M1")), "pristup_txt"] == "ne"
    assert k.loc[str(td.parcel_id("M4")), "pristup_txt"] == "ano"  # přes silnici
    assert k.loc[str(td.parcel_id("S4")), "up_kod"] == "SM"

    wb = openpyxl.load_workbook(tmp_path / "out/kandidati.xlsx")
    assert wb.sheetnames == ["Kandidáti", "Lesní a vyřazené", "Parametry běhu"]
    ws = wb["Kandidáti"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    assert ws.max_row == len(ids) + 1
    hdr = [c.value for c in ws[1]]
    link = ws.cell(row=2, column=hdr.index("Nahlížení do KN") + 1)
    assert link.hyperlink.target.startswith("https://nahlizenidokn.cuzk.gov.cz/ZobrazObjekt.aspx?typ=parcela&id=")

    layers = {n for n, _ in pyogrio.list_layers(tmp_path / "out/kandidati.gpkg")}
    assert layers == {"kandidati", "lesni_vyrazene", "plochy_up_cilove"}
    assert gpd.read_file(tmp_path / "out/kandidati.gpkg", layer="kandidati").crs.to_epsg() == 5514
    html = (tmp_path / "out/mapa.html").read_text(encoding="utf-8")
    assert "ORTOFOTO/MapServer/WMSServer" in html and "Nahlížení do KN" in html
    assert (tmp_path / "out/beh.log").stat().st_size > 0
    assert "CELKEM" in format_summary(res)


def test_selhani_jednoho_ku_neshodi_beh(tmp_path):
    cfg = make_cfg(tmp_path, extra_ku=[{"kod": 999902, "nazev": "Neexistuje"}])
    res = run(cfg, out_dir=tmp_path / "out")
    stav = {r.ku.kod: r.stav for r in res.results}
    assert stav == {td.TEST_KU: "OK", 999902: "CHYBA"}
    assert len(res.kandidati) == len(td.EXPECTED_CANDIDATES)
    assert "CHYBA" in format_summary(res)


def test_chybejici_up_je_srozumitelna_chyba(tmp_path):
    cfg = make_cfg(tmp_path)
    cfg.ku[0].up.cesta = tmp_path / "neni.shp"
    res = run(cfg, out_dir=tmp_path / "out")
    assert not res.ok
    assert "Chybí data ÚP" in res.results[0].zprava
    assert res.kandidati.empty


def test_odkazy():
    assert nahlizeni_url(
        "https://nahlizenidokn.cuzk.gov.cz/ZobrazObjekt.aspx?typ=parcela&id={id}", "2789935604"
    ) == "https://nahlizenidokn.cuzk.gov.cz/ZobrazObjekt.aspx?typ=parcela&id=2789935604"
    assert mapy_url(
        "https://mapy.com/fnc/v1/showmap?mapset=aerial&center={lon},{lat}&zoom=18&marker=true", 16.6, 49.19
    ) == "https://mapy.com/fnc/v1/showmap?mapset=aerial&center=16.600000,49.190000&zoom=18&marker=true"
