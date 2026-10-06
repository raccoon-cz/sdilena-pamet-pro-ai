"""Výstupy: kandidati.xlsx a kandidati.gpkg."""
from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

import geopandas as gpd
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

log = logging.getLogger("parcely")

# (interní sloupec, záhlaví v Excelu, šířka, formát čísla)
KANDIDATI_COLUMNS: list[tuple[str, str, int, str | None]] = [
    ("skore", "Skóre", 8, "0.0"),
    ("obec", "Obec", 18, None),
    ("ku_nazev", "KÚ", 22, None),
    ("cislo", "Číslo parcely", 14, None),
    ("druh_nazev", "Druh pozemku", 22, None),
    ("vymera_filtr", "Výměra (m²)", 12, "#,##0"),
    ("up_plocha", "Plocha v ÚP (kód + název)", 32, None),
    ("podil_pct", "Podíl v ploše (%)", 10, "0.0"),
    ("pristup_txt", "Přístup ke komunikaci", 11, None),
    ("souradnice", "Souřadnice středu (WGS84)", 24, None),
    ("url_nahlizeni", "Nahlížení do KN", 14, None),
    ("url_mapy", "Mapy.cz", 10, None),
    # doplňující sloupce
    ("u_domu_txt", "Sousedí se zastavěnou parcelou (zahrada u domu?)", 14, None),
    ("hranice_zastavena_m", "Společná hranice se zastavěnou parcelou (m)", 13, "0.0"),
    ("vzdalenost_komunikace_m", "Vzdálenost ke komunikaci (m)", 13, "0.0"),
    ("rozpad_skore", "Rozpad skóre", 44, None),
    ("zpusob_nazev", "Způsob využití", 20, None),
    ("up_detail", "Detail plochy ÚP", 22, None),
    ("budova_prekryv_m2", "Max. překryv budovy (m²)", 12, "0.0"),
    ("bud_id_txt", "Budova evidovaná v KN", 11, None),
    ("ku_kod", "Kód KÚ", 9, "0"),
    ("id_parcely", "ID parcely", 14, None),
    ("url_vdp", "RÚIAN (VDP)", 12, None),
]

VYRAZENE_COLUMNS = [
    ("duvod", "Důvod vyřazení", 26, None),
    ("vsechny_duvody", "Všechny důvody", 40, None),
] + KANDIDATI_COLUMNS

LINK_TEXT = {"url_nahlizeni": "Nahlížení", "url_mapy": "Mapy.cz", "url_vdp": "VDP"}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FONT = Font(bold=True, size=12)


def _cell_value(v: Any) -> Any:
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if hasattr(v, "item"):  # numpy skalár
        return v.item()
    return v


def _write_table(ws, df: pd.DataFrame, columns, start_row: int = 1, autofilter: bool = True) -> int:
    cols = list(columns)
    df = df.reindex(columns=[c[0] for c in cols])
    for j, (_, header, width, _) in enumerate(cols, start=1):
        cell = ws.cell(row=start_row, column=j, value=header)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = max(ws.column_dimensions[get_column_letter(j)].width or 0, width)
    ws.row_dimensions[start_row].height = 32
    for i, row in enumerate(df[[c[0] for c in cols]].itertuples(index=False), start=start_row + 1):
        for j, ((key, _, _, fmt), value) in enumerate(zip(cols, row), start=1):
            value = _cell_value(value)
            cell = ws.cell(row=i, column=j)
            if key in LINK_TEXT and value:
                cell.value = LINK_TEXT[key]
                cell.hyperlink = str(value)
                cell.style = "Hyperlink"
            else:
                cell.value = value
                if fmt:
                    cell.number_format = fmt
    last_row = start_row + len(df)
    if autofilter and cols:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(cols))}{max(last_row, start_row)}"
    return last_row


def write_excel(
    path: Path,
    kandidati: pd.DataFrame,
    vyrazene: pd.DataFrame,
    parametry: list[tuple[str, Any]],
    kroky: pd.DataFrame,
    stav_ku: pd.DataFrame,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kandidáti"
    _write_table(ws, kandidati, KANDIDATI_COLUMNS)
    ws.freeze_panes = "A2"

    # „/“ není v názvu listu Excelu povolené, proto „Lesní a vyřazené“.
    ws2 = wb.create_sheet("Lesní a vyřazené")
    _write_table(ws2, vyrazene, VYRAZENE_COLUMNS)
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("Parametry běhu")
    ws3.column_dimensions["A"].width = 38
    ws3.column_dimensions["B"].width = 60
    for i, (k, v) in enumerate(parametry, start=1):
        ws3.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws3.cell(row=i, column=2, value=_cell_value(v)).alignment = Alignment(wrap_text=True)
    row = len(parametry) + 2
    ws3.cell(row=row, column=1, value="Stav katastrálních území").font = SECTION_FONT
    stav_cols = [
        ("ku_kod", "Kód KÚ", 10, "0"), ("ku_nazev", "KÚ", 22, None), ("obec", "Obec", 18, None),
        ("stav", "Stav", 8, None), ("kandidatu", "Kandidátů", 10, "0"), ("budovy", "Zdroj budov", 24, None),
        ("zprava", "Zpráva / poznámky", 60, None),
    ]
    row = _write_table(ws3, stav_ku, stav_cols, start_row=row + 1, autofilter=False) + 2
    ws3.cell(row=row, column=1, value="Počty v krocích filtru").font = SECTION_FONT
    krok_cols = [
        ("ku", "KÚ", 22, None), ("krok", "Krok", 38, None), ("vyrazeno", "Vyřazeno", 10, "#,##0"),
        ("zbyva", "Zbývá", 10, "#,##0"),
    ]
    _write_table(ws3, kroky, krok_cols, start_row=row + 1, autofilter=False)
    for col in "CDEFG":
        ws3.column_dimensions[col].width = max(ws3.column_dimensions[col].width or 0, 12)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    log.info("Excel: %s", path)


def _gpkg_ready(gdf: gpd.GeoDataFrame) -> gpd.GeoDataFrame:
    gdf = gdf.copy()
    for col in gdf.columns:
        if col == gdf.geometry.name:
            continue
        dtype = gdf[col].dtype
        if isinstance(dtype, pd.BooleanDtype) or dtype == bool:
            gdf[col] = gdf[col].map({True: "ano", False: "ne"}).fillna("neznámo")
        elif pd.api.types.is_extension_array_dtype(dtype) and pd.api.types.is_integer_dtype(dtype):
            gdf[col] = gdf[col].astype("float64")
    return gdf


def write_gpkg(path: Path, layers: dict[str, gpd.GeoDataFrame]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        path.unlink()
    for name, gdf in layers.items():
        if gdf is None or gdf.empty:
            continue
        _gpkg_ready(gdf).to_file(path, layer=name, driver="GPKG", engine="pyogrio")
    log.info("GeoPackage: %s", path)
